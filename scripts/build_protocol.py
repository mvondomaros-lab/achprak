"""Build the editable German protocol from the app's task titles and outputs.

Requires python-docx and lxml in a document-authoring environment, not in the app.
Render and visually inspect the DOCX after running this script; see docs/development.md.
"""

from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
STATIC = ROOT / "src/achprak/web/static"
OUTPUT = STATIC / "materials/protokollvorlage.docx"


def build():
    source = html.fromstring((STATIC / "guide.html").read_text())
    tasks = {el.get("id"): el for el in source.xpath('.//details[@class="task"]')}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("192D43")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Title", 25), ("Heading 1", 18), ("Heading 2", 13)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.name, subtitle.font.size = "Arial", Pt(14)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(10)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.core_properties.title = "Molekulare Photoschalter"
    doc.core_properties.subject = "ACh-Pr · TC Versuch · Praktikumsprotokoll"
    doc.core_properties.author = "ACh-Pr"
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "de-DE")
    normal.element.get_or_add_rPr().append(lang)
    footer = sec.footer.paragraphs[0]
    footer.add_run("ACh-Pr · TC Versuch    |    ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("586B80")

    def p(text):
        para = doc.add_paragraph()
        for part in re.split(r"(E(?:trans|cis|TS|max)|ΔE(?:‡|max))", text):
            if part in ("Etrans", "Ecis", "ETS", "Emax"):
                para.add_run("E")
                para.add_run(part[1:]).font.subscript = True
            elif part == "ΔE‡":
                para.add_run("ΔE")
                para.add_run("‡").font.superscript = True
            elif part == "ΔEmax":
                para.add_run("ΔE")
                para.add_run("max").font.subscript = True
            else:
                para.add_run(part)
        return para

    def answer(text="[Ihre Antwort]"):
        para = p(text)
        for run in para.runs:
            run.font.color.rgb = RGBColor.from_string("586B80")

    def task(task_id):
        el = tasks[task_id]
        title = el.find("summary").text.strip()
        doc.add_heading(title, 2)
        text = el.xpath('.//p[@class="protocol-output"]')[0].text_content()
        p(text.removeprefix("Für Ihr Protokoll:").strip())

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for c, name, width in zip(t.rows[0].cells, headers, widths):
            c.text = name
            c.width = Cm(width)
        for row in rows:
            for c, value, width in zip(t.add_row().cells, row, widths):
                c.text = value
                c.width = Cm(width)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            trpr.append(no_split)
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for c in row.cells:
                pr = c._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(
                    qn("w:fill"),
                    "102238" if i == 0 else ("F3F6FA" if i % 2 else "FFFFFF"),
                )
                pr.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                pr.append(margins)
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + side)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                pr.append(borders)
                align = OxmlElement("w:vAlign")
                align.set(qn("w:val"), "center")
                pr.append(align)
                for para in c.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        run.font.size = Pt(9)
                        if i == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        p("")

    doc.add_heading("Molekulare Photoschalter", 0)
    doc.add_paragraph("Struktur, Energie und Absorption im Vergleich", "Subtitle")
    p("ACh-Pr · TC Versuch")
    p("Praktikumsprotokoll")
    p("Name: [Name]    Gruppe: [Gruppe]    Datum: [Datum]")
    p(
        "Dokumentieren Sie Ihre Berechnungen und deren Interpretation. Bearbeiten Sie die Aufgaben in der Webapp. Ersetzen Sie die Angaben in eckigen Klammern und fügen Sie gespeicherte Abbildungen mit Beschriftungen ein. Die Überschriften entsprechen den Aufgaben in der Webapp. Die Felder erweitern sich beim Schreiben."
    )
    p(
        "Speichern Sie die Datei regelmäßig. Laden Sie das ausgefüllte Protokoll in ILIAS hoch. Die Grundlagen finden Sie unter https://mvondomaros-lab.github.io/achprak/theory/."
    )
    doc.add_heading("Strukturen erstellen", 1)
    task("task-compare-configurations")
    answer("Unterschied zwischen cis und trans: [höchstens zwei Sätze]")
    task("task-identify-symmetry")
    answer(
        "Anzahl chemisch verschiedener trans-Azobenzole mit genau einer "
        "Methoxygruppe: [Anzahl]"
    )
    answer(
        "Symmetrieäquivalente Auswahlpositionen: [Fassen Sie die zehn Positionen "
        "in Gruppen zusammen]"
    )
    answer("Begründung anhand der Symmetrie: [höchstens zwei Sätze]")

    doc.add_page_break()
    doc.add_heading("Geometrien und Energiebarrieren untersuchen", 1)
    task("task-compare-optimized-geometries")
    table(
        [
            "Struktur",
            "Elektronische Energie / eV",
            "Diederwinkel / °",
            "Ringabstand / pm",
        ],
        [
            [label, "[Wert]", "[Wert]", "[Wert]"]
            for label in (
                "cis-Startstruktur",
                "trans-Startstruktur",
                "cis-Minimumstruktur",
                "trans-Minimumstruktur",
            )
        ],
        [5, 4, 4, 4],
    )
    answer(
        "ΔE = Etrans − Ecis / (kJ/mol): [Wert]\nGeometrieänderungen und Einordnung des Vorzeichens: [höchstens drei Sätze]"
    )
    task("task-examine-substituent-geometry")
    answer("Vorhersage für die Änderung der Ringstellung: [Ihre Vorhersage]")
    table(
        ["Struktur", "Diederwinkel / °", "Ringabstand / pm"],
        [
            ["unsubstituierte trans-Minimumstruktur", "[Wert]", "[Wert]"],
            ["2,6,2′,6′-Tetramethoxy-Startstruktur", "[Wert]", "[Wert]"],
            ["2,6,2′,6′-Tetramethoxy-Minimumstruktur", "[Wert]", "[Wert]"],
        ],
        [8, 4.5, 4.5],
    )
    answer("Startstruktur: [gleich ausgerichtete 3D-Ansicht einfügen und beschriften]")
    p("\n")
    answer("Minimumstruktur: [gleich ausgerichtete 3D-Ansicht einfügen und beschriften]")
    p("\n")
    answer(
        "Sterischer Einfluss und Vergleich mit unsubstituiertem trans-Azobenzol: [höchstens vier Sätze]"
    )

    doc.add_page_break()
    doc.add_heading("Reaktionspfad und Energiebarrieren", 1)
    task("task-analyze-reaction-path")
    p("")
    answer("[Energieprofil einfügen und beschriften]")
    p("\n\n")
    table(
        ["Richtung", "ΔE‡ / (kJ/mol)", "ΔE‡ / RT bei 298 K"],
        [["trans → cis", "[Wert]", "[Wert]"], ["cis → trans", "[Wert]", "[Wert]"]],
        [5, 6, 6],
    )
    answer(
        "Pfad, Verbindungsprüfung, Barrieren- und RT-Vergleich: [höchstens drei Sätze]"
    )

    doc.add_page_break()
    doc.add_heading("UV/Vis-Spektrum", 1)
    task("task-compare-isomer-spectra")
    table(
        [
            "Minimumstruktur",
            "Maximum / eV",
            "Wellenlänge / nm",
            "UV oder sichtbar",
            "Farbvorhersage",
            "Lichtdurchlässigkeit / %",
        ],
        [
            ["cis", "[Wert]", "[Wert]", "[Bereich]", "[Farbe]", "[Wert]"],
            ["trans", "[Wert]", "[Wert]", "[Bereich]", "[Farbe]", "[Wert]"],
        ],
        [2.5, 2.5, 3.0, 2.7, 3.0, 3.3],
    )
    answer(
        "Bevorzugte Anregung und Grenze der Vorhersage: [höchstens drei Sätze]"
    )
    doc.add_page_break()
    doc.add_heading("Systematische Substituentenreihe", 1)
    task("task-compare-substituent-spectra")
    answer("Zugeordnete Reihe: [4-X, 2-X, 3-X oder 4,4′-X₂]")
    answer(
        "Vorhersage: [Substituent mit der erwarteten stärksten Verschiebung zu kleinerer Energie und kurze Begründung]"
    )
    table(
        [
            "X",
            "Maximum / eV",
            "Wellenlänge / nm",
            "Verschiebung / eV",
            "UV oder sichtbar",
            "Farbvorhersage",
            "Lichtdurchlässigkeit / %",
        ],
        [
            [label, "[Wert]", "[Wert]", "[Wert]", "[Bereich]", "[Farbe]", "[Wert]"]
            for label in ["CH₃", "OCH₃", "N(CH₃)₂", "CF₃", "CN", "NO₂"]
        ],
        [1.8, 2.3, 2.5, 2.1, 2.2, 3.0, 3.1],
    )
    answer("Kleinster Wert von ΔEmax: [Derivat und Wert]")
    answer("Größter Wert von ΔEmax: [Derivat und Wert]")
    answer("Spektrum zum kleinsten Wert von ΔEmax: [Abbildung einfügen und beschriften]")
    p("\n")
    answer("Spektrum zum größten Wert von ΔEmax: [Abbildung einfügen und beschriften]")
    p("\n")
    answer(
        "Trend, Vergleich mit der Vorhersage und Abweichung: [höchstens vier Sätze]"
    )

    doc.add_page_break()
    doc.add_heading("Vergleich der Substituentenreihen", 1)
    task("task-compare-substituent-series")
    table(
        [
            "X",
            "2-X / eV",
            "3-X / eV",
            "4-X / eV",
            "4,4′-X₂ / eV",
        ],
        [
            [label, "[Wert]", "[Wert]", "[Wert]", "[Wert]"]
            for label in ["CH₃", "OCH₃", "N(CH₃)₂", "CF₃", "CN", "NO₂"]
        ],
        [2.0, 3.5, 3.5, 3.5, 4.5],
    )
    answer("Übergreifender Trend 1 mit konkreten Werten: [Trend und Belege]")
    answer("Übergreifender Trend 2 mit konkreten Werten: [Trend und Belege]")
    answer("Abweichung oder begrenzter Gültigkeitsbereich: [Beobachtung]")
    answer("Für den Vergleich vollständiger Spektren ausgewählter Substituent: [X]")
    answer(
        "Vergleich von Absorptionsmaximum, sichtbaren Banden, Farbvorhersage und Lichtdurchlässigkeit: [Ergebnis]"
    )

    doc.add_page_break()
    doc.add_heading("Freie Untersuchung", 1)
    task("task-plan-experiment-series")
    answer(
        "Name und Substitutionsmuster des eigenen Derivats: [Name]\n"
        "Bezug zur systematischen Reihe und Begründung der Auswahl: [Ergebnis und Begründung]\n"
        "Vorhersage: [Ihre Vorhersage vor den Rechnungen]"
    )
    table(
        [
            "Konfiguration",
            "Maximum / eV",
            "Wellenlänge / nm",
            "Farbvorhersage",
            "Lichtdurchlässigkeit / %",
        ],
        [
            ["cis", "[Wert]", "[Wert]", "[Farbe]", "[Wert]"],
            ["trans", "[Wert]", "[Wert]", "[Farbe]", "[Wert]"],
        ],
        [3.0, 3.0, 3.5, 3.5, 4.0],
    )
    answer("cis: [Spektrum einfügen und beschriften]")
    p("\n\n")
    answer("trans: [Spektrum einfügen und beschriften]")
    p("\n\n")
    answer(
        "Vergleich der vollständigen Spektren und Prüfung der Vorhersage: [Stützt das Ergebnis den Trend, widerspricht es ihm oder ist keine eindeutige Aussage möglich?]"
    )
    answer("Längstwelliges Gruppenresultat: [Name und Wellenlänge]")
    answer(
        "Stärkster vorhergesagter cis/trans-Farbunterschied: [Name und Farben]"
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
